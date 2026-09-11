#!/usr/bin/env python3
"""Inspect legacy CBSRMT workbooks and emit a machine-readable archaeology report.

The inspector never maps workbook columns into canonical catalog entities. Its
purpose is discovery: enumerate sheets, dimensions, headers, value types,
blank rates, distinct counts, likely keys, formulas, merged cells, and samples.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import statistics
import sys
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def classify(value: Any) -> str:
    if value is None or value == "":
        return "blank"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, datetime):
        return "datetime"
    if isinstance(value, date):
        return "date"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return "number"
    if isinstance(value, str) and value.startswith("="):
        return "formula"
    if isinstance(value, str):
        return "text"
    return "unknown"


def normalize_header(value: Any, column_number: int) -> str:
    if value is None or str(value).strip() == "":
        return f"column_{column_number}"
    return " ".join(str(value).strip().split())


def serializable(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return str(value)
    return value


def inspect_matrix(name: str, rows: list[list[Any]], *, hidden: bool = False) -> dict[str, Any]:
    row_count = len(rows)
    column_count = max((len(row) for row in rows), default=0)
    padded = [row + [None] * (column_count - len(row)) for row in rows]

    header_row = 1 if row_count else None
    headers = [normalize_header(v, i + 1) for i, v in enumerate(padded[0])] if row_count else []
    data_rows = padded[1:] if row_count else []

    columns = []
    for index in range(column_count):
        values = [row[index] for row in data_rows]
        nonblank = [v for v in values if v not in (None, "")]
        type_counts = Counter(classify(v) for v in values)
        textified = [json.dumps(serializable(v), sort_keys=True, default=str) for v in nonblank]
        distinct_count = len(set(textified))
        blank_count = len(values) - len(nonblank)
        candidate_key = bool(nonblank) and len(nonblank) == len(values) and distinct_count == len(nonblank)

        date_values = [v for v in nonblank if isinstance(v, (date, datetime))]
        numeric_values = [float(v) for v in nonblank if isinstance(v, (int, float)) and not isinstance(v, bool)]

        summary: dict[str, Any] = {
            "columnNumber": index + 1,
            "header": headers[index] if index < len(headers) else f"column_{index + 1}",
            "blankCount": blank_count,
            "blankPercent": round((blank_count / len(values) * 100), 2) if values else 0.0,
            "distinctCount": distinct_count,
            "candidateKey": candidate_key,
            "typeCounts": dict(sorted(type_counts.items())),
            "samples": [serializable(v) for v in nonblank[:5]],
        }
        if date_values:
            summary["dateRange"] = {
                "min": min(date_values).isoformat(),
                "max": max(date_values).isoformat(),
            }
        if numeric_values:
            summary["numericRange"] = {"min": min(numeric_values), "max": max(numeric_values)}
            if len(numeric_values) >= 2:
                summary["numericMedian"] = statistics.median(numeric_values)
        columns.append(summary)

    return {
        "sheetName": name,
        "hidden": hidden,
        "rowCount": row_count,
        "columnCount": column_count,
        "headerRow": header_row,
        "headers": headers,
        "columns": columns,
        "sampleRows": [[serializable(v) for v in row] for row in data_rows[:10]],
    }


def inspect_xlsx(path: Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=False, data_only=False)
    reports = []
    for worksheet in workbook.worksheets:
        rows = [[cell.value for cell in row] for row in worksheet.iter_rows()]
        report = inspect_matrix(worksheet.title, rows, hidden=worksheet.sheet_state != "visible")
        report["mergedRanges"] = [str(rng) for rng in worksheet.merged_cells.ranges]
        report["formulaCount"] = sum(
            1 for row in worksheet.iter_rows() for cell in row if cell.data_type == "f"
        )
        reports.append(report)
    return reports


def _xls_cell_value(cell: Any, datemode: int, xlrd_module: Any) -> Any:
    if cell.ctype == xlrd_module.XL_CELL_EMPTY:
        return None
    if cell.ctype == xlrd_module.XL_CELL_DATE:
        converted = xlrd_module.xldate_as_datetime(cell.value, datemode)
        if converted.time().isoformat() == "00:00:00":
            return converted.date()
        return converted
    if cell.ctype == xlrd_module.XL_CELL_BOOLEAN:
        return bool(cell.value)
    if cell.ctype == xlrd_module.XL_CELL_ERROR:
        return f"#XLERROR:{int(cell.value)}"
    return cell.value


def inspect_xls(path: Path) -> list[dict[str, Any]]:
    try:
        import xlrd
    except ImportError as exc:
        raise RuntimeError("xlrd is required to inspect .xls files") from exc

    workbook = xlrd.open_workbook(path, formatting_info=False)
    reports = []
    for sheet in workbook.sheets():
        rows = [
            [_xls_cell_value(sheet.cell(row_index, column_index), workbook.datemode, xlrd)
             for column_index in range(sheet.ncols)]
            for row_index in range(sheet.nrows)
        ]
        report = inspect_matrix(
            sheet.name,
            rows,
            hidden=bool(getattr(sheet, "visibility", 0)),
        )
        report["mergedRanges"] = []
        report["formulaCount"] = None
        reports.append(report)
    return reports


def inspect(path: Path) -> dict[str, Any]:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        sheets = inspect_xlsx(path)
        workbook_format = suffix[1:]
    elif suffix == ".xls":
        sheets = inspect_xls(path)
        workbook_format = "xls"
    else:
        raise ValueError(f"unsupported workbook format: {suffix}")

    return {
        "schemaVersion": "1.0",
        "fileName": path.name,
        "fileFormat": workbook_format,
        "sha256": sha256(path),
        "sheetCount": len(sheets),
        "sheets": sheets,
        "mappingStatus": "REVIEW_REQUIRED",
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args(argv)

    if not args.workbook.exists():
        print(f"error: workbook not found: {args.workbook}", file=sys.stderr)
        return 2

    try:
        report = inspect(args.workbook)
    except Exception as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 1

    rendered = json.dumps(report, indent=2, ensure_ascii=False)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(rendered + "\n", encoding="utf-8")
        print(f"WROTE: {args.output}")
    else:
        print(rendered)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
