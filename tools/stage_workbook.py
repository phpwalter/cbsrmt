#!/usr/bin/env python3
"""Stage workbook rows and cells without canonicalizing them."""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Any

if __package__ in (None, ""):
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import yaml

from tools.inspect_workbook import _xls_cell_value, inspect, normalize_header

IMPORTER_VERSION = "workbook-stager/0.1.0"
SOURCE_TYPE = "legacy_workbook"


def checksum(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def serialize_value(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def read_xlsx(path: Path) -> list[tuple[str, bool, list[list[Any]]]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=False, data_only=False)
    return [
        (
            ws.title,
            ws.sheet_state != "visible",
            [[cell.value for cell in row] for row in ws.iter_rows()],
        )
        for ws in workbook.worksheets
    ]


def read_xls(path: Path) -> list[tuple[str, bool, list[list[Any]]]]:
    try:
        import xlrd
    except ImportError as exc:
        raise RuntimeError("xlrd is required to stage .xls files") from exc

    workbook = xlrd.open_workbook(path, formatting_info=False)
    result: list[tuple[str, bool, list[list[Any]]]] = []
    for sheet in workbook.sheets():
        rows = [
            [_xls_cell_value(sheet.cell(row_index, column_index), workbook.datemode, xlrd)
             for column_index in range(sheet.ncols)]
            for row_index in range(sheet.nrows)
        ]
        result.append((sheet.name, bool(getattr(sheet, "visibility", 0)), rows))
    return result


def workbook_rows(path: Path):
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return read_xlsx(path)
    if suffix == ".xls":
        return read_xls(path)
    raise ValueError(f"unsupported workbook format: {suffix}")


def ensure_source(cur, path: Path) -> str:
    cur.execute(
        """
        INSERT INTO provenance.sources (source_type, name, description, uri)
        VALUES (%s, %s, %s, %s)
        ON CONFLICT (source_type, name)
        DO UPDATE SET description = EXCLUDED.description
        RETURNING source_id::text
        """,
        (SOURCE_TYPE, path.name, "Legacy CBSRMT workbook", path.name),
    )
    return cur.fetchone()[0]


def ensure_batch(cur, source_id: str, file_checksum: str) -> tuple[str, str]:
    cur.execute(
        """
        INSERT INTO provenance.import_batches
            (source_id, source_checksum, importer_version, status)
        VALUES (%s::uuid, %s, %s, 'running')
        ON CONFLICT (source_id, source_checksum, importer_version)
        DO UPDATE SET source_checksum = EXCLUDED.source_checksum
        RETURNING import_batch_id::text, status
        """,
        (source_id, file_checksum, IMPORTER_VERSION),
    )
    row = cur.fetchone()
    return row[0], row[1]


def load_mapping(path: Path | None) -> dict[str, Any] | None:
    if path is None:
        return None
    return yaml.safe_load(path.read_text(encoding="utf-8"))


def assert_mapping_allows_canonicalization(mapping: dict[str, Any] | None) -> None:
    if mapping is None:
        raise RuntimeError("canonicalization requires a reviewed mapping manifest")
    status = mapping.get("workbook", {}).get("status")
    if status != "APPROVED":
        raise RuntimeError(f"canonicalization blocked: workbook mapping status is {status!r}")


def existing_stage_summary(cur, batch_id: str, file_checksum: str) -> dict[str, Any] | None:
    cur.execute(
        """
        SELECT w.workbook_id::text,
               count(DISTINCT s.workbook_sheet_id) AS sheet_count,
               count(DISTINCT r.workbook_row_id) AS row_count,
               count(c.workbook_cell_id) AS cell_count
        FROM staging.workbooks w
        LEFT JOIN staging.workbook_sheets s ON s.workbook_id = w.workbook_id
        LEFT JOIN staging.workbook_rows r ON r.workbook_sheet_id = s.workbook_sheet_id
        LEFT JOIN staging.workbook_cells c ON c.workbook_row_id = r.workbook_row_id
        WHERE w.import_batch_id = %s::uuid
          AND w.file_checksum = %s
        GROUP BY w.workbook_id
        """,
        (batch_id, file_checksum),
    )
    row = cur.fetchone()
    if row is None:
        return None
    return {
        "workbookId": row[0],
        "sheetCount": int(row[1]),
        "rowCount": int(row[2]),
        "cellCount": int(row[3]),
    }


def stage(database_url: str, workbook_path: Path, mapping_path: Path | None = None) -> dict[str, Any]:
    import psycopg

    report = inspect(workbook_path)
    file_checksum = checksum(workbook_path)
    sheets = workbook_rows(workbook_path)
    mapping = load_mapping(mapping_path)

    with psycopg.connect(database_url) as conn:
        with conn.cursor() as cur:
            source_id = ensure_source(cur, workbook_path)
            batch_id, batch_status = ensure_batch(cur, source_id, file_checksum)

            existing = existing_stage_summary(cur, batch_id, file_checksum)
            if existing is not None:
                if batch_status != "completed":
                    raise RuntimeError(
                        "staging evidence already exists for this workbook checksum but the import batch is not completed"
                    )
                conn.rollback()
                return {
                    **existing,
                    "importBatchId": batch_id,
                    "mappingStatus": (mapping or {}).get("workbook", {}).get("status", "UNMAPPED"),
                    "canonicalization": "BLOCKED",
                    "stagingStatus": "UNCHANGED",
                }

            if batch_status == "completed":
                raise RuntimeError(
                    "completed workbook staging batch exists without staging evidence; refusing to recreate immutable evidence"
                )

            cur.execute(
                """
                INSERT INTO staging.workbooks
                    (import_batch_id, source_id, file_name, file_checksum, workbook_format, sheet_count)
                VALUES (%s::uuid, %s::uuid, %s, %s, %s, %s)
                RETURNING workbook_id::text
                """,
                (batch_id, source_id, workbook_path.name, file_checksum, report["fileFormat"], len(sheets)),
            )
            workbook_id = cur.fetchone()[0]

            row_total = 0
            cell_total = 0
            for sheet_index, (sheet_name, hidden, rows) in enumerate(sheets):
                row_count = len(rows)
                column_count = max((len(row) for row in rows), default=0)
                header_row = 1 if row_count else None
                cur.execute(
                    """
                    INSERT INTO staging.workbook_sheets
                        (workbook_id, sheet_index, sheet_name, row_count, column_count, header_row, hidden)
                    VALUES (%s::uuid, %s, %s, %s, %s, %s, %s)
                    RETURNING workbook_sheet_id::text
                    """,
                    (workbook_id, sheet_index, sheet_name, row_count, column_count, header_row, hidden),
                )
                sheet_id = cur.fetchone()[0]

                headers = [normalize_header(v, i + 1) for i, v in enumerate(rows[0])] if rows else []
                for row_number, row in enumerate(rows, start=1):
                    payload = {
                        (headers[i] if i < len(headers) else f"column_{i + 1}"): serialize_value(value)
                        for i, value in enumerate(row)
                    }
                    cur.execute(
                        """
                        INSERT INTO staging.workbook_rows
                            (workbook_sheet_id, row_number, row_payload)
                        VALUES (%s::uuid, %s, %s::jsonb)
                        RETURNING workbook_row_id::text
                        """,
                        (sheet_id, row_number, json.dumps(payload, ensure_ascii=False)),
                    )
                    row_id = cur.fetchone()[0]
                    row_total += 1

                    for column_number, value in enumerate(row, start=1):
                        column_letter = _column_letter(column_number)
                        header_name = headers[column_number - 1] if column_number <= len(headers) else f"column_{column_number}"
                        value_type = _value_type(value)
                        cur.execute(
                            """
                            INSERT INTO staging.workbook_cells
                                (workbook_row_id, column_number, column_letter, cell_address,
                                 header_name, raw_value, normalized_value, value_type, formula)
                            VALUES (%s::uuid, %s, %s, %s, %s, %s, %s, %s, %s)
                            """,
                            (
                                row_id,
                                column_number,
                                column_letter,
                                f"{column_letter}{row_number}",
                                header_name,
                                serialize_value(value),
                                serialize_value(value),
                                value_type,
                                str(value) if value_type == "formula" else None,
                            ),
                        )
                        cell_total += 1

            cur.execute(
                """
                UPDATE provenance.import_batches
                SET status = 'completed', completed_at = now(),
                    rows_seen = %s, rows_accepted = %s, rows_rejected = 0
                WHERE import_batch_id = %s::uuid
                """,
                (row_total, row_total, batch_id),
            )
        conn.commit()

    return {
        "workbookId": workbook_id,
        "importBatchId": batch_id,
        "sheetCount": len(sheets),
        "rowCount": row_total,
        "cellCount": cell_total,
        "mappingStatus": (mapping or {}).get("workbook", {}).get("status", "UNMAPPED"),
        "canonicalization": "BLOCKED",
        "stagingStatus": "CREATED",
    }


def _column_letter(index: int) -> str:
    letters = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def _value_type(value: Any) -> str:
    if value is None or value == "":
        return "blank"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, datetime):
        return "datetime"
    if isinstance(value, date):
        return "date"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return "number"
    if isinstance(value, str) and value.startswith("="):
        return "formula"
    if isinstance(value, str):
        return "text"
    return "unknown"


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--database-url", default=os.getenv("DATABASE_URL"))
    parser.add_argument("--mapping", type=Path)
    args = parser.parse_args(argv)

    if not args.workbook.exists():
        print(f"error: workbook not found: {args.workbook}", file=sys.stderr)
        return 2
    if not args.database_url:
        print("error: --database-url or DATABASE_URL is required", file=sys.stderr)
        return 2

    try:
        result = stage(args.database_url, args.workbook, args.mapping)
    except Exception as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 1

    print(json.dumps(result, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
